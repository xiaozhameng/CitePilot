"""手动上传的知识库：解析 → 切分 → 落盘 → 可增删。

────────────────────────────────────────────────────────────
为什么需要它
────────────────────────────────────────────────────────────
内置语料是一本固定的电子书。真实使用里，用户手里一定有别的资料：
自己写的规范、某产品的接口文档、一份 PDF 报告。
没有上传能力的知识库，等于只能回答"开发者预先想到的那点问题"。

────────────────────────────────────────────────────────────
存储布局
────────────────────────────────────────────────────────────
    data/kb/
      documents.json           文档清单（谁在库里、各有多少块）
      <doc_id>/
        original.<ext>         上传的原件
        chunks.jsonl           切分结果

为什么保留原件？
因为切分策略一定会改（换 embedding 模型就要重新调块长）。
原件留着，重跑只是一条命令；不留，就得让用户重新上传一遍。

────────────────────────────────────────────────────────────
为什么不做增量索引（而是一上传就全量重建）
────────────────────────────────────────────────────────────
BM25 的 IDF 是**全局统计量** —— 加一份文档会改变所有词的权重，
所以"增量更新"本质上必须重算全部已存文档的分数，工作量等于全量重建。
在此之上再加脏标记 / 分段索引 / 分数合并，在几百块的规模上是纯负债。
真正需要增量的阈值在万级以上，那时该换的是存储（pgvector），不是打补丁。

另外：embedding 缓存是按「单条文本的 sha1」做的，所以重建时
原有文档的向量全部命中缓存，只有新增的块要真正调一次模型 ——
这让"每次都全量重建"在成本上也站得住。
"""

from __future__ import annotations

import io
import json
import re
import shutil
import threading
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterator

from .config import SETTINGS
from .ingest import load_chunks, save_chunks
from .schema import Chunk
from .split import BlockPacker

# 支持的上传格式。加新格式只需在这里登记 + 写一个抽取函数。
SUPPORTED_SUFFIXES = {
    ".md", ".markdown", ".txt", ".text",
    ".html", ".htm", ".pdf", ".docx",
}

# 抽取函数的统一形状：bytes -> 一串 (level, title, body)
# level=0 表示正文块（只有 body），>=1 表示标题块
Blocks = Iterator[tuple[int, str, str]]


def _decode(data: bytes) -> str:
    """字节转文本。中文文档可能是 GBK/GB18030，所以不能只试 utf-8。"""
    for enc in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    # 全都不行就用替换字符兜底 —— 宁可有个别乱码，也好过整个文件失败
    return data.decode("utf-8", errors="replace")


# ------------------------------------------------------------------
# 各格式的文本抽取
# ------------------------------------------------------------------
_MD_HEAD = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def _blocks_markdown(data: bytes) -> Blocks:
    text = _decode(data)
    buf: list[str] = []
    in_code = False
    for raw in text.splitlines():
        line = raw.rstrip()
        if _MD_FENCE.match(line):
            in_code = not in_code
            if line.strip():
                buf.append(line.strip())
            continue
        if not in_code:
            m = _MD_HEAD.match(line)
            if m:
                if buf:
                    yield (0, "", "\n".join(buf))
                    buf = []
                yield (len(m.group(1)), m.group(2).strip(), "")
                continue
        if line.strip():
            buf.append(line.strip())
        elif buf:
            yield (0, "", "\n".join(buf))
            buf = []
    if buf:
        yield (0, "", "\n".join(buf))


def _blocks_plain(data: bytes) -> Blocks:
    """纯文本：按空行分段。没有标题结构可用，出处粒度只能是文档级。"""
    for para in re.split(r"\n\s*\n", _decode(data)):
        para = para.strip()
        if para:
            yield (0, "", para)


def _blocks_html(data: bytes) -> Blocks:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "lxml")
    # 只挑有语义的标签，且 find_all 按文档顺序返回 —— 顺序不能乱，
    # 乱了就会出现"引用的小节标题和正文对不上"。
    for tag in soup.find_all(
        ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th", "blockquote", "pre"]
    ):
        text = tag.get_text(" ", strip=True)
        if not text:
            continue
        name = tag.name or ""
        if name.startswith("h"):
            try:
                yield (int(name[1]), text, "")
            except ValueError:
                yield (0, "", text)
        else:
            yield (0, "", text)


def _blocks_pdf(data: bytes) -> Blocks:
    """PDF：逐页抽文本，每页当作一个"小节"。

    PDF 里没有可靠的标题结构（视觉上的标题和正文抽出来都是普通文本行），
    但「第 N 页」本身就是一个**有用的出处粒度** —— 比"这个文档的某处"强得多，
    而且和纸质材料对得上号。这是对 PDF 能做的最诚实的取舍。
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    for i, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:
            # 单页抽失败不该毁掉整份文档
            continue
        text = text.strip()
        if text:
            yield (2, f"第 {i} 页", text)


def _blocks_docx(data: bytes) -> Blocks:
    import docx

    d = docx.Document(io.BytesIO(data))
    for p in d.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = ""
        try:
            style = (p.style.name if p.style is not None else "") or ""
        except Exception:
            style = ""
        m = re.match(r"Heading\s*(\d+)", style, re.I)
        if m:
            yield (int(m.group(1)), text, "")
        elif style.lower().startswith(("title", "标题")):
            yield (1, text, "")
        else:
            yield (0, "", text)
    # 表格不在 paragraphs 里，要单独取。表格往往是最密集的知识点。
    for tbl in d.tables:
        rows: list[str] = []
        for row in tbl.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            yield (0, "", "\n".join(rows))


_EXTRACTORS: dict[str, Callable[[bytes], Blocks]] = {
    ".md": _blocks_markdown,
    ".markdown": _blocks_markdown,
    ".txt": _blocks_plain,
    ".text": _blocks_plain,
    ".html": _blocks_html,
    ".htm": _blocks_html,
    ".pdf": _blocks_pdf,
    ".docx": _blocks_docx,
}


def extract_blocks(suffix: str, data: bytes) -> Blocks:
    """按扩展名分发到对应的抽取器。加新格式只改这一处和上面的字典。"""
    fn = _EXTRACTORS.get(suffix.lower())
    if fn is None:
        raise ValueError(
            f"不支持的文件类型「{suffix or '无扩展名'}」。"
            f"目前支持：{'、'.join(sorted(SUPPORTED_SUFFIXES))}"
        )
    return fn(data)


# ------------------------------------------------------------------
# 文档元信息与清单
# ------------------------------------------------------------------
@dataclass
class DocumentInfo:
    doc_id: str
    name: str          # 显示名（文件名去掉扩展名）
    filename: str      # 上传时的原始文件名
    suffix: str
    size: int
    chunk_count: int
    added_at: str

    def to_dict(self) -> dict:
        return asdict(self)


_ID_BAD = re.compile(r'[/\\:*?"<>|\s\x00-\x1f]+')


def make_doc_id(filename: str, existing: set[str]) -> str:
    """从文件名生成 doc_id。

    为什么允许中文？
    因为 doc_id 会出现在 URL 和引用标签里。用哈希虽然"安全"，但用户看到
    `doc-a3f9c2` 根本不知道是哪份文件，排查问题时反而更费劲。
    中文在 URL 里做百分号编码即可，代价远小于可读性的收益。
    """
    stem = Path(filename).stem
    base = _ID_BAD.sub("-", stem).strip("-._")[:48] or "doc"
    if base not in existing:
        return base
    for i in range(2, 200):
        cand = f"{base}-{i}"
        if cand not in existing:
            return cand
    return f"{base}-{uuid.uuid4().hex[:6]}"


class KnowledgeBase:
    """上传知识库的读写入口。所有方法都能被多个请求并发调用。"""

    def __init__(self, root: Path | None = None) -> None:
        self.root = Path(root or SETTINGS.kb_dir)
        self.root.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()

    @property
    def manifest_path(self) -> Path:
        return self.root / "documents.json"

    # ---- 清单 ----
    def list_documents(self) -> list[DocumentInfo]:
        if not self.manifest_path.exists():
            return []
        try:
            raw = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        except Exception:
            # 清单坏了不让整个知识库不可用 —— 顶多显示不出列表
            return []
        out: list[DocumentInfo] = []
        for d in raw if isinstance(raw, list) else []:
            if not isinstance(d, dict):
                continue
            fields = {k: v for k, v in d.items() if k in DocumentInfo.__dataclass_fields__}
            try:
                out.append(DocumentInfo(**fields))
            except TypeError:
                continue
        return sorted(out, key=lambda x: x.added_at)

    def _write_manifest(self, docs: list[DocumentInfo]) -> None:
        # 先写临时文件再 rename：中途崩了也不会留下半个损坏的 JSON
        tmp = self.manifest_path.with_name(self.manifest_path.name + ".tmp")
        tmp.write_text(
            json.dumps([d.to_dict() for d in docs], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        tmp.replace(self.manifest_path)

    # ---- 增删 ----
    def add_document(self, filename: str, data: bytes) -> DocumentInfo:
        suffix = Path(filename).suffix.lower()
        if suffix not in _EXTRACTORS:
            raise ValueError(
                f"不支持的文件类型「{suffix or '无扩展名'}」。"
                f"目前支持：{'、'.join(sorted(SUPPORTED_SUFFIXES))}"
            )
        if not data:
            raise ValueError("文件是空的")

        with self._lock:
            doc_id = make_doc_id(filename, {d.doc_id for d in self.list_documents()})
            doc_dir = self.root / doc_id
            doc_dir.mkdir(parents=True, exist_ok=True)
            (doc_dir / f"original{suffix}").write_bytes(data)

            try:
                packer = BlockPacker(doc_id=doc_id, doc_name=Path(filename).stem)
                for level, title, body in extract_blocks(suffix, data):
                    if level > 0:
                        packer.heading(level, title)
                    if body:
                        packer.body(body)
                chunks = packer.result()
            except Exception as exc:  # noqa: BLE001
                shutil.rmtree(doc_dir, ignore_errors=True)
                raise ValueError(f"解析「{filename}」失败：{type(exc).__name__}: {exc}") from exc

            if not chunks:
                # 切出 0 块必须当失败处理，这点很重要：
                # 否则库里会留下一条"显示已上传、但检索永远召不回"的记录，
                # 用户提问时怎么都得不到引用，这类问题最难排查。
                # 最常见的原因是扫描版 PDF（整页都是图片，没有文字层）。
                shutil.rmtree(doc_dir, ignore_errors=True)
                raise ValueError(
                    f"「{filename}」里没有切出任何文字内容。"
                    "如果是扫描版 PDF（整页都是图片），需要先做 OCR 才能入库。"
                )

            save_chunks(chunks, doc_dir / "chunks.jsonl")
            info = DocumentInfo(
                doc_id=doc_id,
                name=Path(filename).stem,
                filename=filename,
                suffix=suffix,
                size=len(data),
                chunk_count=len(chunks),
                added_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            )
            docs = [d for d in self.list_documents() if d.doc_id != doc_id]
            docs.append(info)
            self._write_manifest(docs)
            return info

    def remove_document(self, doc_id: str) -> bool:
        with self._lock:
            docs = self.list_documents()
            keep = [d for d in docs if d.doc_id != doc_id]
            if len(keep) == len(docs):
                return False
            shutil.rmtree(self.root / doc_id, ignore_errors=True)
            self._write_manifest(keep)
            return True

    # ---- 读取 ----
    def load_chunks(self) -> list[Chunk]:
        """读出库里所有文档的块。

        单份文档坏了只跳过它，不让整个知识库不可用 ——
        一份文件解析出问题，不该连带把其他文档的问答能力一起弄没。
        """
        out: list[Chunk] = []
        for info in self.list_documents():
            path = self.root / info.doc_id / "chunks.jsonl"
            if not path.exists():
                continue
            try:
                out.extend(load_chunks(path))
            except Exception:
                continue
        return out

    def raw_chunk(self, doc_id: str, chunk_id: str) -> dict | None:
        """给引用深链用的原文接口：返回某一块的完整文本与出处。

        这是「引用可追溯」对上传文档同样成立的唯一方式 ——
        上传的文档没有可点的网页，用户点引用只能路由到这个接口看原文。
        """
        path = self.root / doc_id / "chunks.jsonl"
        if not path.exists():
            return None
        try:
            chunks = load_chunks(path)
        except Exception:
            return None
        for c in chunks:
            if c.chunk_id == chunk_id:
                return {
                    "doc_id": c.doc_id,
                    "doc_name": c.doc_name,
                    "section": c.section_title,
                    "label": c.label,
                    "text": c.text,
                }
        return None
